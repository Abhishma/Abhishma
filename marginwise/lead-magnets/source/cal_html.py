import sys, html
from openpyxl import load_workbook
src, out = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
E = lambda s: html.escape(str(s or ""))
BR = lambda s: E(s).replace("\n", "<br>")

cal = wb["Calendar"]
heads = [c.value for c in cal[1]]
ix = {h: i for i, h in enumerate(heads)}
rows = [r for r in cal.iter_rows(min_row=2, values_only=True) if r[0]]

weeks = {}
for r in rows:
    weeks.setdefault(r[ix["Week"]], []).append(r)

def steps_html(s):
    items = [l.split(". ", 1)[1] if ". " in l[:4] else l for l in str(s).split("\n") if l.strip()]
    return "<ol class='steps'>" + "".join(f"<li>{E(i)}</li>" for i in items) + "</ol>"

cards = []
for w, rs in weeks.items():
    d0 = rs[0][ix["Date"]].strftime("%d %b"); d1 = rs[-1][ix["Date"]].strftime("%d %b")
    cards.append(f"<h3 class='week' id='week{w}'>Week {w} <span>{d0} to {d1}</span></h3>")
    for n, r in enumerate(rs):
        d = r[ix["Date"]]
        cid = f"cap-{w}-{n}"
        fmt = E(r[ix["Format"]])
        kind = "carousel" if "Carousel" in fmt else "reel"
        flags = []
        if r[ix["Status"]] == "Posted": flags.append("<span class='chip'>Posted</span>")
        if r[ix["Pin?"]] == "Yes": flags.append("<span class='chip pin'>Pin</span>")
        if r[ix["Trial Reel first?"]] == "Yes": flags.append("<span class='chip trial'>Trial Reel first</span>")
        cards.append(f"""
<article class="post">
  <header class="post-head">
    <div class="when"><b>{d.strftime('%a %d %b')}</b><span class="mono">{E(r[ix['IST']])} IST · {E(r[ix['Sydney']])} Sydney · {E(r[ix['London']])} London · {E(r[ix['New York']])} New York</span></div>
    <div class="tags"><span class="chip {kind}">{fmt}</span><span class="chip">{E(r[ix['Market']])}</span><span class="chip">{E(r[ix['Pillar']])}</span>{''.join(flags)}</div>
  </header>
  <p class="hook">{E(r[ix['Hook (on screen)']])}</p>
  <div class="block"><h4>What it shows</h4><p>{E(r[ix['Outline']])}</p></div>
  <div class="block"><h4>Step by step (no face)</h4>{steps_html(r[ix['Step by step (no face)']])}</div>
  <div class="block"><div class="caphead"><h4>Caption</h4><button type="button" class="copy" data-target="{cid}">Copy caption</button></div>
    <p class="caption" id="{cid}">{BR(r[ix['Caption']])}</p></div>
  <dl class="meta"><div><dt>Comment keyword</dt><dd class="mono">{E(r[ix['CTA keyword']])}</dd></div>
  <div><dt>Sends</dt><dd>{E(r[ix['Lead magnet']])}</dd></div><div><dt>Record in</dt><dd>{E(r[ix['Recording batch']])}</dd></div></dl>
</article>""")

def table(ws, cols=None):
    rs = list(ws.iter_rows(values_only=True))
    h = rs[0]; body = rs[1:]
    keep = cols or list(range(len(h)))
    out = "<div class='tw'><table><thead><tr>" + "".join(f"<th>{E(h[i])}</th>" for i in keep) + "</tr></thead><tbody>"
    for r in body:
        if not any(r): continue
        cells = []
        for i in keep:
            v = r[i]
            if hasattr(v, "strftime"): v = v.strftime("%a %d %b")
            cells.append(f"<td>{E(v)}</td>")
        out += "<tr>" + "".join(cells) + "</tr>"
    return out + "</tbody></table></div>"

weeknav = "".join(f'<a href="#week{w}">Wk {w}</a>' for w in weeks)
sop = table(wb["How to make (no face)"])
batches = table(wb["Recording batches"])
stories = table(wb["Stories"])
setup_rows = [r for r in wb["Setup before 14 Oct"].iter_rows(min_row=2, values_only=True) if r[1]]
setup = "<ul class='checks'>" + "".join(
    f"<li><input type='checkbox' id='s{i}'><label for='s{i}'><b>{E(t)}</b><span>{E(why)}</span></label></li>"
    for i, (_, t, why) in enumerate(setup_rows)) + "</ul>"

page = f"""<title>Marginwise Content Calendar</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:opsz,wght@12..96,600;12..96,700&family=Figtree:wght@400;500;600&family=JetBrains+Mono:wght@400;500&display=swap">
<style>
/* Layout: one reading column of post cards grouped by week, sticky section nav on top. */
:root {{
  --bg:#f3f4f7; --surface:#ffffff; --ink:#18203a; --muted:#5b6478; --line:#dde1ea;
  --accent:#1f6f5c; --accent-soft:#e3f1ec; --reel:#e8ecf7; --carousel:#f6efe1; --pin:#fbe7e3; --pin-ink:#8a2f1b;
  --display:"Bricolage Grotesque", "Segoe UI", system-ui, sans-serif;
  --body:"Figtree", "Segoe UI", system-ui, sans-serif;
  --mono:"JetBrains Mono", ui-monospace, Menlo, monospace;
}}
@media (prefers-color-scheme: dark) {{ :root:not([data-theme="light"]) {{
  --bg:#10131c; --surface:#181d2a; --ink:#e7eaf2; --muted:#9aa3b8; --line:#2a3142;
  --accent:#5cc4a6; --accent-soft:#16332c; --reel:#222a3d; --carousel:#33291a; --pin:#3a201b; --pin-ink:#f2a493; color-scheme:dark }} }}
:root[data-theme="dark"] {{
  --bg:#10131c; --surface:#181d2a; --ink:#e7eaf2; --muted:#9aa3b8; --line:#2a3142;
  --accent:#5cc4a6; --accent-soft:#16332c; --reel:#222a3d; --carousel:#33291a; --pin:#3a201b; --pin-ink:#f2a493; color-scheme:dark }}
* {{ box-sizing:border-box }}
body {{ background:var(--bg); color:var(--ink); font-family:var(--body); font-size:15px; line-height:1.55 }}
.wrap {{ max-width:760px; margin:0 auto; padding-inline:16px; padding-block:8px 64px }}
nav {{ position:sticky; top:env(safe-area-inset-top,0px); z-index:5; background:var(--bg); border-bottom:1px solid var(--line); }}
nav .in {{ max-width:760px; margin:0 auto; padding:10px 16px; display:flex; gap:8px; overflow-x:auto; scrollbar-width:none }}
nav a {{ flex:0 0 auto; text-decoration:none; color:var(--ink); font-weight:600; font-size:13px; padding:7px 12px; border:1px solid var(--line); border-radius:999px; background:var(--surface) }}
nav a:focus-visible, button:focus-visible {{ outline:2px solid var(--accent); outline-offset:2px }}
h1 {{ font-family:var(--display); font-size:clamp(26px,6vw,36px); line-height:1.1; margin:24px 0 6px; text-wrap:balance }}
.lede {{ color:var(--muted); margin:0 0 8px; max-width:62ch }}
.rule {{ font-size:13px; color:var(--muted); margin:0 0 4px }}
h2 {{ font-family:var(--display); font-size:22px; margin:40px 0 12px; scroll-margin-top:64px }}
h3.week {{ font-family:var(--display); font-size:18px; margin:28px 0 10px; scroll-margin-top:64px }}
h3.week span {{ font-family:var(--mono); font-size:12px; color:var(--muted); font-weight:400; margin-left:6px }}
.post {{ background:var(--surface); border:1px solid var(--line); border-radius:10px; padding:16px; margin-bottom:14px; display:grid; gap:12px }}
.post-head {{ display:flex; flex-wrap:wrap; justify-content:space-between; gap:8px }}
.when {{ display:grid; gap:2px }} .when b {{ font-size:15px }}
.mono {{ font-family:var(--mono); font-size:12px; color:var(--muted); font-variant-numeric:tabular-nums }}
.tags {{ display:flex; flex-wrap:wrap; gap:6px; align-items:flex-start }}
.chip {{ font-size:12px; padding:2px 8px; border-radius:6px; background:var(--bg); border:1px solid var(--line); white-space:nowrap }}
.chip.reel {{ background:var(--reel) }} .chip.carousel {{ background:var(--carousel) }}
.chip.pin {{ background:var(--pin); color:var(--pin-ink); border-color:transparent; font-weight:600 }}
.chip.trial {{ background:var(--accent-soft); color:var(--accent); border-color:transparent; font-weight:600 }}
.hook {{ font-family:var(--display); font-size:19px; font-weight:600; line-height:1.3; margin:0; text-wrap:balance }}
.block h4 {{ margin:0 0 4px; font-size:11px; letter-spacing:.08em; text-transform:uppercase; color:var(--muted) }}
.block p {{ margin:0 }}
.steps {{ margin:0; padding-left:20px; display:grid; gap:6px }}
.caphead {{ display:flex; justify-content:space-between; align-items:center; gap:8px; margin-bottom:4px }}
.caphead h4 {{ margin:0 }}
.caption {{ background:var(--bg); border-radius:8px; padding:10px 12px; font-size:14px; user-select:text }}
button.copy {{ font:600 12px var(--body); color:var(--accent); background:var(--accent-soft); border:0; border-radius:6px; padding:6px 10px; cursor:pointer }}
.meta {{ display:flex; flex-wrap:wrap; gap:8px 20px; margin:0; padding-top:10px; border-top:1px solid var(--line) }}
.meta div {{ min-width:0 }} .meta dt {{ font-size:11px; color:var(--muted) }} .meta dd {{ margin:0; font-size:13px; font-weight:500 }}
.meta dd.mono {{ color:var(--accent); font-weight:500; font-size:13px }}
.tw {{ overflow-x:auto; background:var(--surface); border:1px solid var(--line); border-radius:10px }}
table {{ border-collapse:collapse; width:100%; font-size:14px; min-width:520px }}
th, td {{ text-align:left; vertical-align:top; padding:10px 12px; border-bottom:1px solid var(--line) }}
th {{ font-size:12px; color:var(--muted); font-weight:600 }}
tr:last-child td {{ border-bottom:0 }}
ul.checks {{ list-style:none; padding:0; margin:0; display:grid; gap:8px }}
ul.checks li {{ display:flex; gap:10px; background:var(--surface); border:1px solid var(--line); border-radius:10px; padding:12px }}
ul.checks input {{ margin-top:4px; accent-color:var(--accent); width:18px; height:18px; flex:0 0 auto }}
ul.checks label {{ display:grid; gap:2px; min-width:0 }} ul.checks span {{ color:var(--muted); font-size:13px }}
@media (prefers-reduced-motion: no-preference) {{ html {{ scroll-behavior:smooth }} }}
</style>

<nav aria-label="Sections"><div class="in">
  <a href="#calendar">Calendar</a>{weeknav}
  <a href="#how">How to make</a><a href="#batches">Recording</a><a href="#stories">Stories</a><a href="#setup">Setup</a>
</div></nav>
<main class="wrap">
  <h1>@marginwise.in content calendar</h1>
  <p class="lede">14 Oct to 13 Dec 2026. Australia, US and UK, adding Germany and Canada from 17 Nov. No face on camera: every post shows the work, with your voiceover.</p>
  <p class="rule">From 18 Oct: Australia posts at 13:30 IST; US/UK and all-market posts at 19:30 IST. Reply to every comment in the first hour. Don't publish a keyword until its auto-DM works.</p>
  <h2 id="calendar">Calendar</h2>
  {''.join(cards)}
  <h2 id="how">How to make (no face)</h2>{sop}
  <h2 id="batches">Recording batches</h2>{batches}
  <h2 id="stories">Stories, every week</h2>{stories}
  <h2 id="setup">Setup before the first post</h2>{setup}
</main>
<script>
document.querySelectorAll("button.copy").forEach(function (b) {{
  b.addEventListener("click", function () {{
    var el = document.getElementById(b.dataset.target);
    var text = el.innerText;
    function selectIt() {{ var r = document.createRange(); r.selectNodeContents(el); var s = window.getSelection(); s.removeAllRanges(); s.addRange(r); b.textContent = "Selected, copy now"; }}
    try {{
      navigator.clipboard.writeText(text).then(function () {{ b.textContent = "Copied"; setTimeout(function () {{ b.textContent = "Copy caption"; }}, 1600); }}, selectIt);
    }} catch (e) {{ selectIt(); }}
  }});
}});
document.querySelectorAll("ul.checks input").forEach(function (cb) {{
  try {{ cb.checked = localStorage.getItem("mw-" + cb.id) === "1"; }} catch (e) {{}}
  cb.addEventListener("change", function () {{ try {{ localStorage.setItem("mw-" + cb.id, cb.checked ? "1" : "0"); }} catch (e) {{}} }});
}});
</script>
"""
open(out, "w").write(page)
print("ok", len(rows))
