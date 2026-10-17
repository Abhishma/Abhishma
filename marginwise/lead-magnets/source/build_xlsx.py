"""Builds the three Marginwise lead-magnet workbooks."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]

FONT = "Arial"
BLUE = "0000FF"
NAVY = "1F2A44"
INPUT_FILL = PatternFill("solid", start_color="FFF6CC")
HEAD_FILL = PatternFill("solid", start_color="1F2A44")
SECTION_FILL = PatternFill("solid", start_color="E8ECF4")
RESULT_FILL = PatternFill("solid", start_color="E6F2EA")
thin = Side(style="thin", color="C9CED8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

USD = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
INR = '"₹"#,##0;("₹"#,##0);-'
INR2 = '"₹"#,##0.00;("₹"#,##0.00);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.00"x"'
NUM = '#,##0;(#,##0);-'


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 16)
    ws["A2"] = sub
    ws["A2"].font = f(False, "555555", 10, italic=True)


def section(ws, row, text, cols=4):
    for c in range(1, cols + 1):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    ws.cell(row=row, column=1, value=text).font = f(True, NAVY, 11)


def inp(cell, value, fmt):
    cell.value = value
    cell.font = f(color=BLUE)
    cell.fill = INPUT_FILL
    cell.number_format = fmt
    cell.border = BOX
    cell.alignment = Alignment(horizontal="right")


def calc(cell, formula, fmt, bold=False, fill=None):
    cell.value = formula
    cell.font = f(bold)
    cell.number_format = fmt
    cell.border = BOX
    cell.alignment = Alignment(horizontal="right")
    if fill:
        cell.fill = fill


def label(ws, row, text, bold=False, col=1):
    c = ws.cell(row=row, column=col, value=text)
    c.font = f(bold)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    return c


def note(ws, row, text, col):
    c = ws.cell(row=row, column=col, value=text)
    c.font = f(False, "555555", 9, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="center")


def readme(ws, heading, lines):
    ws.column_dimensions["A"].width = 110
    ws["A1"] = heading
    ws["A1"].font = f(True, NAVY, 16)
    r = 3
    for kind, text in lines:
        c = ws.cell(row=r, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h":
            c.font = f(True, NAVY, 11)
            r += 1
        elif kind == "input":
            c.font = f(color=BLUE)
            c.fill = INPUT_FILL
        else:
            c.font = f()
        r += 1
    ws.sheet_view.showGridLines = False


def base_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False


FOOT = ("Marginwise | marginwise.in | Built for D2C and Shopify brands in India, the US and the UK. "
        "Scenarios to plan with, not a promise of results.")


# ---------------------------------------------------------------- USD workbook
def build_usd(path):
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    readme(rm, "Marginwise Profit Calculator (USD)", [
        ("h", "What this does"),
        ("t", "Tab 'ROAS to Profit': turns the ROAS your ad platform reports into what you actually keep after "
              "over-attribution, returns, product cost, shipping and payment fees."),
        ("t", "Tab 'Conversion Lift': shows what a 10%, 20% or 30% conversion lift is worth in gross margin, "
              "with the same traffic and no extra ad spend."),
        ("h", "How to use"),
        ("input", "Edit only the cells shaded like this (blue text on yellow). Everything else calculates."),
        ("t", "The figures already filled in are an example store. Replace them with your last full month."),
        ("t", "Black numbers are formulas. Do not type over them."),
        ("h", "Where the defaults come from"),
        ("t", "Over-attribution of 25%: vendor estimates put platform over-counting at 20-40% (Luca, citing the "
              "Ecom CFO 2026 report). Not independently verified. Replace it with your own incrementality test or "
              "post-purchase survey data if you have it."),
        ("t", "Conversion lift example: matches the example store on marginwise.in (40,000 sessions, 2.0% "
              "conversion, $75 AOV, 60% gross margin)."),
        ("t", "All other defaults are illustrative, not benchmarks."),
        ("h", "Want a second pair of eyes?"),
        ("t", "Book a free 30-minute call at marginwise.in. You leave with your top three leaks and a written summary."),
        ("t", FOOT),
    ])

    ws = wb.create_sheet("ROAS to Profit")
    base_widths(ws, {"A": 52, "B": 16, "C": 70})
    title(ws, "ROAS to Profit", "Monthly view. Edit the yellow cells only.")
    section(ws, 4, "INPUTS", 3)
    rows = [
        (5, "Monthly ad spend ($)", 20000, USD, ""),
        (6, "ROAS reported by the ad platform", 3.0, MULT, "From Meta or Google Ads Manager."),
        (7, "Over-attribution (% of platform revenue not truly caused by ads)", 0.25, PCT,
         "Vendor estimates: 20-40%. Use your own data if you have it."),
        (8, "Average order value ($, after discounts)", 60, USD2, ""),
        (9, "Return and refund rate (% of revenue)", 0.08, PCT, ""),
        (10, "Product cost, COGS (% of net revenue)", 0.35, PCT, ""),
        (11, "Shipping and fulfilment cost per order ($)", 8, USD2, "Pick, pack and postage."),
        (12, "Payment processing (% of revenue)", 0.029, PCT, "Check your Shopify Payments or Stripe rate."),
        (13, "Payment processing fixed fee per order ($)", 0.30, USD2, ""),
    ]
    for r, lab, val, fmt, nt in rows:
        label(ws, r, lab)
        inp(ws.cell(row=r, column=2), val, fmt)
        note(ws, r, nt, 3)

    section(ws, 15, "RESULTS (MONTHLY)", 3)
    res = [
        (16, "Revenue reported by the ad platform", "=B5*B6", USD, False),
        (17, "Revenue actually driven by ads", "=B16*(1-B7)", USD, False),
        (18, "Orders driven by ads", "=IFERROR(B17/B8,0)", NUM, False),
        (19, "Less returns and refunds", "=-B17*B9", USD, False),
        (20, "Net revenue", "=B17+B19", USD, True),
        (21, "Less product cost (COGS)", "=-B20*B10", USD, False),
        (22, "Less shipping and fulfilment", "=-B18*B11", USD, False),
        (23, "Less payment processing", "=-(B17*B12+B18*B13)", USD, False),
        (24, "Contribution before ad spend", "=B20+B21+B22+B23", USD, True),
        (25, "Less ad spend", "=-B5", USD, False),
        (26, "Contribution after ad spend", "=B24+B25", USD, True),
    ]
    for r, lab, fm, fmt, b in res:
        label(ws, r, lab, b)
        calc(ws.cell(row=r, column=2), fm, fmt, b)

    section(ws, 28, "THE NUMBERS THAT MATTER", 3)
    label(ws, 29, "True ROAS (net revenue / ad spend)", True)
    calc(ws["B29"], "=IFERROR(B20/B5,0)", MULT, True, RESULT_FILL)
    note(ws, 29, "Compare this with the platform number in B6.", 3)
    label(ws, 30, "Contribution per order after ads ($)", True)
    calc(ws["B30"], "=IFERROR(B26/B18,0)", USD2, True, RESULT_FILL)
    label(ws, 31, "Contribution after ads (% of net revenue)", True)
    calc(ws["B31"], "=IFERROR(B26/B20,0)", PCT, True, RESULT_FILL)
    label(ws, 32, "Break-even platform ROAS", True)
    calc(ws["B32"],
         '=IF(((1-B9)*(1-B10)-B11/B8-B12-B13/B8)<=0,"Not reachable",'
         '1/(((1-B9)*(1-B10)-B11/B8-B12-B13/B8)*(1-B7)))', MULT, True, RESULT_FILL)
    note(ws, 32, "The ROAS your platform must report before ads stop losing money. "
                 "'Not reachable' means costs exceed revenue on every order.", 3)
    label(ws, 33, "Verdict", True)
    calc(ws["B33"],
         '=IF(B26>=0,"Profitable after ads",IF(ISNUMBER(B32),"Losing money: needs "&TEXT(B32,"0.00")&"x",'
         '"Losing money on every order"))', "@", True, RESULT_FILL)
    ws["B33"].alignment = Alignment(horizontal="left", wrap_text=True)
    ws.row_dimensions[33].height = 30
    note(ws, 35, FOOT, 1)

    ws2 = wb.create_sheet("Conversion Lift")
    base_widths(ws2, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 40})
    title(ws2, "What a conversion lift is worth", "Same traffic, no extra ad spend. Edit the yellow cells only.")
    section(ws2, 4, "YOUR STORE TODAY", 5)
    for r, lab, val, fmt in [
        (5, "Sessions per month", 40000, NUM),
        (6, "Conversion rate", 0.02, '0.00%'),
        (7, "Average order value ($)", 75, USD2),
        (8, "Gross margin (%)", 0.60, PCT),
    ]:
        label(ws2, r, lab)
        inp(ws2.cell(row=r, column=2), val, fmt)
    label(ws2, 9, "Orders per month")
    calc(ws2["B9"], "=B5*B6", NUM)
    label(ws2, 10, "Revenue per month", True)
    calc(ws2["B10"], "=B9*B7", USD, True)

    section(ws2, 12, "SCENARIOS", 5)
    label(ws2, 13, "Conversion lift you want to model", True)
    for col, v in zip("BCD", [0.10, 0.20, 0.30]):
        inp(ws2[f"{col}13"], v, '+0%')
    scen = [
        (14, "New conversion rate", "=$B$6*(1+{c}13)", '0.00%', False),
        (15, "Extra orders per month", "=$B$5*$B$6*{c}13", NUM, False),
        (16, "Extra revenue per month", "={c}15*$B$7", USD, True),
        (17, "Extra gross margin per month", "={c}16*$B$8", USD, False),
        (18, "Extra gross margin per year", "={c}17*12", USD, True),
    ]
    for r, lab, fm, fmt, b in scen:
        label(ws2, r, lab, b)
        for col in "BCD":
            calc(ws2[f"{col}{r}"], fm.format(c=col), fmt, b,
                 RESULT_FILL if r == 18 else None)
    note(ws2, 20, "Lifts are scenarios to plan with, not a promise of results. For context: Baymard Institute "
                  "puts average cart abandonment at 70.22% across 50 studies, and estimates the average large "
                  "ecommerce site can gain 35% higher conversion through better checkout design.", 1)
    ws2.merge_cells("A20:E20")
    ws2.row_dimensions[20].height = 42
    note(ws2, 22, FOOT, 1)
    wb.save(path)


# ---------------------------------------------------------------- INR workbook
def build_inr(path):
    wb = Workbook()
    rm = wb.active
    rm.title = "Read me"
    readme(rm, "Marginwise Contribution Margin Calculator (INR)", [
        ("h", "What this does"),
        ("t", "Tab 'Margin per Order': shows what one order actually earns you once COD, RTO, shipping, "
              "ad cost, gateway fees and GST are counted. It compares your current COD mix with a scenario "
              "where more buyers pay upfront."),
        ("t", "Tab 'GST on Ads': shows the cash your ad budget really needs once 18% GST is added."),
        ("h", "How to use"),
        ("input", "Edit only the cells shaded like this (blue text on yellow). Everything else calculates."),
        ("t", "Column B is your store today. Column C is a what-if: change the COD share and prepaid discount "
              "to see whether pushing prepaid is worth it."),
        ("t", "Enter costs excluding GST where you claim input tax credit on them."),
        ("h", "Where the defaults come from"),
        ("t", "RTO rates of 26% on COD and 2% on prepaid: Shipway ShipNotes FY25, reported by MediaBrief "
              "(26% RTO on COD orders vs under 2% on prepaid)."),
        ("t", "Prices, costs and shares are an illustrative example store, not benchmarks. Replace them with "
              "your last full month."),
        ("t", "The scenario assumes your conversion rate does not change when you add a prepaid discount. "
              "Test it before rolling it out."),
        ("h", "Want a second pair of eyes?"),
        ("t", "Book a free 30-minute call at marginwise.in. You leave with your top three leaks and a written summary."),
        ("t", FOOT),
    ])

    ws = wb.create_sheet("Margin per Order")
    base_widths(ws, {"A": 52, "B": 16, "C": 16, "D": 60})
    title(ws, "Contribution margin per order", "Edit the yellow cells only. Column C is a what-if scenario.")
    for col, txt in (("B", "Today"), ("C", "Scenario")):
        c = ws[f"{col}4"]
        c.value = txt
        c.font = f(True, "FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    ws["A4"].fill = HEAD_FILL
    ws["D4"].fill = HEAD_FILL
    ws["A4"].value = "INPUTS"
    ws["A4"].font = f(True, "FFFFFF")

    inputs = [
        (5, "Price paid by customer (₹, incl. GST)", 1000, 1000, INR, ""),
        (6, "GST rate on the product", 0.18, 0.18, PCT, "Use your HSN rate (often 5%, 12% or 18%)."),
        (7, "Product cost per order, COGS (₹)", 300, 300, INR, ""),
        (8, "Ad cost per order placed (₹)", 250, 250, INR, "Ad spend / orders placed, including orders that later RTO."),
        (9, "Forward shipping per order (₹)", 70, 70, INR, ""),
        (10, "Packaging per order (₹)", 20, 20, INR, ""),
        (11, "Return shipping per RTO (₹)", 70, 70, INR, ""),
        (12, "COD share of orders", 0.60, 0.35, PCT, "Scenario: fewer COD orders after a prepaid offer."),
        (13, "RTO rate on COD orders", 0.26, 0.26, PCT, "Shipway FY25: 26%."),
        (14, "RTO rate on prepaid orders", 0.02, 0.02, PCT, "Shipway FY25: under 2%."),
        (15, "COD handling fee per delivered COD order (₹)", 35, 35, INR, "From your courier or aggregator rate card."),
        (16, "Payment gateway fee on prepaid (% of price)", 0.02, 0.02, PCT, ""),
        (17, "Prepaid discount offered (% of price)", 0, 0.05, PCT, "Applies to prepaid orders only."),
        (18, "Stock lost per RTO (% of COGS damaged or unsellable)", 0.10, 0.10, PCT, ""),
        (19, "Orders placed per month", 3000, 3000, NUM, ""),
    ]
    for r, lab, a, b, fmt, nt in inputs:
        label(ws, r, lab)
        inp(ws[f"B{r}"], a, fmt)
        inp(ws[f"C{r}"], b, fmt)
        note(ws, r, nt, 4)

    section(ws, 21, "PER ORDER PLACED", 4)
    calcs = [
        (22, "Share of orders delivered", "={c}12*(1-{c}13)+(1-{c}12)*(1-{c}14)", PCT, False),
        (23, "Share of orders returned (RTO)", "=1-{c}22", PCT, False),
        (24, "Net revenue, excluding GST",
         "=({c}12*(1-{c}13)*{c}5+(1-{c}12)*(1-{c}14)*{c}5*(1-{c}17))/(1+{c}6)", INR2, True),
        (25, "Less product cost of delivered orders", "=-{c}22*{c}7", INR2, False),
        (26, "Less stock lost on RTO", "=-{c}23*{c}7*{c}18", INR2, False),
        (27, "Less ad cost", "=-{c}8", INR2, False),
        (28, "Less forward shipping and packaging", "=-({c}9+{c}10)", INR2, False),
        (29, "Less return shipping on RTO", "=-{c}23*{c}11", INR2, False),
        (30, "Less COD handling fees", "=-{c}12*(1-{c}13)*{c}15", INR2, False),
        (31, "Less payment gateway fees", "=-(1-{c}12)*(1-{c}14)*{c}5*(1-{c}17)*{c}16", INR2, False),
        (32, "Contribution per order placed", "=SUM({c}24:{c}31)", INR2, True),
    ]
    for r, lab, fm, fmt, b in calcs:
        label(ws, r, lab, b)
        for col in "BC":
            calc(ws[f"{col}{r}"], fm.format(c=col), fmt, b)

    section(ws, 34, "THE NUMBERS THAT MATTER", 4)
    key = [
        (35, "Contribution per delivered order", "=IFERROR({c}32/{c}22,0)", INR2),
        (36, "Contribution margin (% of net revenue)", "=IFERROR({c}32/{c}24,0)", PCT),
        (37, "Ad and shipping spend lost on RTO, per order placed", "={c}23*({c}8+{c}9+{c}10+{c}11)", INR2),
        (38, "Contribution per month", "={c}32*{c}19", INR),
    ]
    for r, lab, fm, fmt in key:
        label(ws, r, lab, True)
        for col in "BC":
            calc(ws[f"{col}{r}"], fm.format(c=col), fmt, True, RESULT_FILL)
    label(ws, 39, "Scenario vs today, per month", True)
    calc(ws["C39"], "=C38-B38", INR, True, RESULT_FILL)
    note(ws, 39, "Positive means the scenario earns more per month.", 4)
    note(ws, 41, FOOT, 1)
    ws.merge_cells("A41:D41")

    g = wb.create_sheet("GST on Ads")
    base_widths(g, {"A": 52, "B": 18, "C": 70})
    title(g, "GST on your ad budget", "Not tax advice. Confirm the treatment with your CA.")
    section(g, 4, "INPUTS", 3)
    label(g, 5, "Monthly ad budget (₹)")
    inp(g["B5"], 100000, INR)
    label(g, 6, "GST rate on advertising")
    inp(g["B6"], 0.18, PCT)
    note(g, 6, "Advertising services are taxed at 18% GST.", 3)
    label(g, 7, "Billed by a foreign entity? (Yes or No)")
    inp(g["B7"], "Yes", "@")
    note(g, 7, "Yes if the invoice is from Meta Platforms Ireland or Google Ireland (no GSTIN on the invoice).", 3)
    label(g, 8, "Are you GST-registered and able to claim input credit? (Yes or No)")
    inp(g["B8"], "Yes", "@")
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    g.add_data_validation(dv)
    dv.add("B7")
    dv.add("B8")

    section(g, 10, "RESULTS", 3)
    label(g, 11, "GST on the budget")
    calc(g["B11"], "=B5*B6", INR)
    label(g, 12, "How it is paid")
    calc(g["B12"], '=IF(B7="Yes","Reverse charge: you self-assess and pay it","Charged on the invoice")', "@")
    g["B12"].alignment = Alignment(horizontal="left", wrap_text=True)
    g.row_dimensions[12].height = 30
    label(g, 13, "Cash that leaves your account this month", True)
    calc(g["B13"], "=B5+B11", INR, True, RESULT_FILL)
    label(g, 14, "Recoverable later as input tax credit")
    calc(g["B14"], '=IF(B8="Yes",B11,0)', INR)
    label(g, 15, "Final cost after credit", True)
    calc(g["B15"], "=B13-B14", INR, True, RESULT_FILL)
    label(g, 16, "Cash tied up until the credit is claimed", True)
    calc(g["B16"], "=B14", INR, True, RESULT_FILL)
    note(g, 18, "Under reverse charge, the IGST is paid in cash first and then claimed as credit, so plan "
                "working capital for the full amount. Source: growthlife.in, 'At what monthly ad spend does "
                "hiring a marketing agency actually start making sense in India?' (Sep 2026).", 1)
    g.merge_cells("A18:C18")
    g.row_dimensions[18].height = 42
    note(g, 20, FOOT, 1)
    g.merge_cells("A20:C20")
    wb.save(path)


# ---------------------------------------------------------------- RTO workbook
STATES = [
    "Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh", "Goa", "Gujarat", "Haryana",
    "Himachal Pradesh", "Jharkhand", "Karnataka", "Kerala", "Madhya Pradesh", "Maharashtra", "Manipur",
    "Meghalaya", "Mizoram", "Nagaland", "Odisha", "Punjab", "Rajasthan", "Sikkim", "Tamil Nadu", "Telangana",
    "Tripura", "Uttar Pradesh", "Uttarakhand", "West Bengal", "Andaman and Nicobar Islands", "Chandigarh",
    "Dadra and Nagar Haveli and Daman and Diu", "Delhi", "Jammu and Kashmir", "Ladakh", "Lakshadweep",
    "Puducherry",
]
MAXR = 20001  # order rows 2..20001
PIN_ROWS = 2000


def build_rto(path):
    wb = Workbook()
    rm = wb.active
    rm.title = "How to use"
    readme(rm, "Marginwise RTO Finder: which states and pin codes are costing you", [
        ("h", "What this does"),
        ("t", "Ranks states and pin codes by RTO rate and by what RTO costs you, so you can switch off COD "
              "only where it loses money instead of everywhere."),
        ("h", "Steps (about 10 minutes)"),
        ("t", "1. Export the last 90 days of shipped orders from Shopify, WooCommerce or your shipping "
              "aggregator (Shiprocket, Shipway, etc.)."),
        ("t", "2. Paste them into the 'Orders' tab from row 2, matching the column order: Order ID, Order "
              "date, Pin code, State, Payment method, Final status, Order value."),
        ("t", "3. Payment method must be COD or Prepaid. Final status must be Delivered or RTO. Leave orders "
              "still in transit or cancelled before dispatch out of the paste, or mark them anything else "
              "and they will be ignored."),
        ("t", "4. State names must match the list on the 'By State' tab (for example 'Tamil Nadu', not 'TN')."),
        ("t", "5. For pin codes: copy the Pin code column, remove duplicates (Google Sheets: Data > Data "
              "cleanup > Remove duplicates; Excel: Data > Remove Duplicates), and paste the unique list "
              "into column A of 'By Pin Code'."),
        ("input", "6. Check the yellow cells on the 'Settings' tab (shipping, ad cost per order, thresholds)."),
        ("h", "Example row for the Orders tab"),
        ("t", "#10234 | 2026-09-14 | 302001 | Rajasthan | COD | RTO | 1299"),
        ("h", "How to read the result"),
        ("t", "'Restrict COD or verify' means COD RTO in that area is at or above your threshold and there "
              "are enough orders to trust the rate. Options: prepaid-only, a small prepaid discount, or "
              "confirming COD orders by WhatsApp or call before dispatch."),
        ("t", "'Too few orders' means the sample is too small to judge. Do not switch COD off on a handful "
              "of orders."),
        ("t", "Benchmark: Shipway ShipNotes FY25 reports 26% RTO on COD orders in India versus under 2% on "
              "prepaid (via MediaBrief)."),
        ("h", "Want a second pair of eyes?"),
        ("t", "Book a free 30-minute call at marginwise.in. Bring this sheet and we will go through it."),
        ("t", FOOT),
    ])

    st = wb.create_sheet("Settings")
    base_widths(st, {"A": 50, "B": 16, "C": 60})
    title(st, "Settings", "Edit the yellow cells only.")
    sets = [
        (4, "Forward shipping per order (₹)", 70, INR, ""),
        (5, "Packaging per order (₹)", 20, INR, ""),
        (6, "Return shipping per RTO (₹)", 70, INR, ""),
        (7, "Ad cost per order placed (₹)", 250, INR, "Ad spend / orders placed. Wasted on every RTO."),
        (8, "Minimum COD orders before judging an area", 20, NUM, "Below this, the rate is too noisy to act on."),
        (9, "COD RTO rate that triggers action", 0.30, PCT, "Set this above your store average."),
    ]
    for r, lab, v, fmt, nt in sets:
        label(st, r, lab)
        inp(st[f"B{r}"], v, fmt)
        note(st, r, nt, 3)
    label(st, 11, "Cost of one COD RTO (₹)", True)
    calc(st["B11"], "=B4+B5+B6+B7", INR, True, RESULT_FILL)
    note(st, 11, "Forward + packaging + return + wasted ad cost. Excludes damaged stock.", 3)

    od = wb.create_sheet("Orders")
    heads = ["Order ID", "Order date", "Pin code", "State", "Payment method", "Final status", "Order value (₹)"]
    for i, h in enumerate(heads, 1):
        c = od.cell(row=1, column=i, value=h)
        c.font = f(True, "FFFFFF")
        c.fill = HEAD_FILL
    base_widths(od, {"A": 14, "B": 14, "C": 12, "D": 30, "E": 18, "F": 16, "G": 16})
    od.sheet_view.showGridLines = True
    od.freeze_panes = "A2"
    dvp = DataValidation(type="list", formula1='"COD,Prepaid"', allow_blank=True)
    dvs = DataValidation(type="list", formula1='"Delivered,RTO"', allow_blank=True)
    od.add_data_validation(dvp)
    od.add_data_validation(dvs)
    dvp.add(f"E2:E{MAXR}")
    dvs.add(f"F2:F{MAXR}")
    od["I1"] = "Paste from row 2. Payment method: COD or Prepaid. Final status: Delivered or RTO."
    od["I1"].font = f(False, "555555", 9, True)

    rngP = f"Orders!$C$2:$C${MAXR}"
    rngS = f"Orders!$D$2:$D${MAXR}"
    rngM = f"Orders!$E$2:$E${MAXR}"
    rngF = f"Orders!$F$2:$F${MAXR}"

    def summary_sheet(name, key_head, keys, key_range, n_rows, blank_guard):
        ws = wb.create_sheet(name)
        base_widths(ws, {"A": 34, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 18,
                         "I": 10, "J": 26})
        title(ws, f"RTO {name.lower()}", "Calculated from the Orders tab. Do not edit except column A on By Pin Code.")
        heads = [key_head, "COD shipped", "COD RTO", "COD RTO rate", "Prepaid shipped", "Prepaid RTO",
                 "Prepaid RTO rate", "Cost of COD RTO (₹)", "Cost rank", "Action"]
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=4, column=i, value=h)
            c.font = f(True, "FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "B5"
        first, last = 5, 4 + n_rows
        for i in range(n_rows):
            r = first + i
            a = ws.cell(row=r, column=1)
            if keys is not None:
                a.value = keys[i]
                a.font = f()
            else:
                a.font = f(color=BLUE)
                a.fill = INPUT_FILL
            g = (lambda body: f'=IF($A{r}="","",{body})') if blank_guard else (lambda body: "=" + body)
            k = f"{key_range},$A{r}"
            fx = {
                "B": g(f'COUNTIFS({k},{rngM},"COD",{rngF},"Delivered")+COUNTIFS({k},{rngM},"COD",{rngF},"RTO")'),
                "C": g(f'COUNTIFS({k},{rngM},"COD",{rngF},"RTO")'),
                "D": g(f"IFERROR(C{r}/B{r},0)"),
                "E": g(f'COUNTIFS({k},{rngM},"Prepaid",{rngF},"Delivered")+COUNTIFS({k},{rngM},"Prepaid",{rngF},"RTO")'),
                "F": g(f'COUNTIFS({k},{rngM},"Prepaid",{rngF},"RTO")'),
                "G": g(f"IFERROR(F{r}/E{r},0)"),
                "H": g(f"C{r}*Settings!$B$11"),
                "I": g(f'IF(N(H{r})=0,"",RANK(H{r},$H${first}:$H${last}))'),
                "J": g(f'IF(B{r}=0,"No COD orders",IF(B{r}<Settings!$B$8,"Too few orders",'
                       f'IF(D{r}>=Settings!$B$9,"Restrict COD or verify","Keep COD")))'),
            }
            fmts = {"B": NUM, "C": NUM, "D": PCT, "E": NUM, "F": NUM, "G": PCT, "H": INR, "I": "0", "J": "@"}
            for col, fm in fx.items():
                c = ws[f"{col}{r}"]
                c.value = fm
                c.font = f()
                c.number_format = fmts[col]
                c.alignment = Alignment(horizontal="left" if col == "J" else "right")
        tr = last + 1
        ws.cell(row=tr, column=1, value="All").font = f(True)
        for col in "BCEFH":
            c = ws[f"{col}{tr}"]
            c.value = f"=SUM({col}{first}:{col}{last})"
            c.font = f(True)
            c.number_format = INR if col == "H" else NUM
        ws[f"D{tr}"] = f"=IFERROR(C{tr}/B{tr},0)"
        ws[f"G{tr}"] = f"=IFERROR(F{tr}/E{tr},0)"
        for col in "DG":
            ws[f"{col}{tr}"].font = f(True)
            ws[f"{col}{tr}"].number_format = PCT
        for col in "ABCDEFGHIJ":
            ws[f"{col}{tr}"].fill = RESULT_FILL
        ws.auto_filter.ref = f"A4:J{last}"
        return ws

    summary_sheet("By State", "State", STATES, rngS, len(STATES), False)
    ps = summary_sheet("By Pin Code", "Pin code (paste unique list)", None, rngP, PIN_ROWS, True)
    ps["A3"] = "Paste your unique pin codes into the yellow column A from row 5. Use the filter on 'Cost rank' to sort."
    ps["A3"].font = f(False, "555555", 9, True)
    wb.save(path)


build_usd(f"{OUT}/marginwise-profit-calculator-usd.xlsx")
build_inr(f"{OUT}/marginwise-margin-calculator-inr.xlsx")
build_rto(f"{OUT}/marginwise-rto-finder.xlsx")
print("built")
